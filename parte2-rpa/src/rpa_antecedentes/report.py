"""Generación del reporte Excel y del resumen en consola."""

from __future__ import annotations

import contextlib
import logging
import os
import tempfile
from collections import Counter
from collections.abc import Sequence
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from rpa_antecedentes.models import QueryResult, Status

log = logging.getLogger(__name__)

#: Columnas exigidas por la prueba, en este orden.
HEADERS = ("cedula", "nombres", "estado", "detalle", "fecha_consulta")
DATETIME_FORMAT = "yyyy-mm-dd hh:mm:ss"

_HEADER_FILL = PatternFill("solid", fgColor="1F4E8C")
_STATUS_FILL = {
    Status.CLEAN: PatternFill("solid", fgColor="E3F3E4"),
    Status.WITH_RECORDS: PatternFill("solid", fgColor="FFF1D6"),
    Status.NOT_FOUND: PatternFill("solid", fgColor="E9EEF3"),
    Status.ERROR: PatternFill("solid", fgColor="FDE4E8"),
}


class ReportWriteError(Exception):
    """No se pudo escribir el reporte (carpeta sin permisos, archivo abierto en Excel...)."""


@dataclass(frozen=True, slots=True)
class RunSummary:
    results: int
    clean: int
    with_records: int
    not_found: int
    errors: int
    rows_read: int
    blank_rows: int
    duplicates_removed: int
    invalid: int
    started_at: datetime
    finished_at: datetime
    service_url: str

    @property
    def total(self) -> int:
        """Filas del reporte = cédulas únicas procesadas."""
        return self.results

    @property
    def duration_seconds(self) -> float:
        return (self.finished_at - self.started_at).total_seconds()


def build_summary(
    results: Sequence[QueryResult],
    *,
    rows_read: int,
    blank_rows: int,
    duplicates_removed: int,
    started_at: datetime,
    finished_at: datetime,
    service_url: str,
) -> RunSummary:
    counts = Counter(r.status for r in results)
    invalid = sum(1 for r in results if r.status is Status.ERROR and r.queried_at is None)
    return RunSummary(
        results=len(results),
        clean=counts[Status.CLEAN],
        with_records=counts[Status.WITH_RECORDS],
        not_found=counts[Status.NOT_FOUND],
        errors=counts[Status.ERROR],
        rows_read=rows_read,
        blank_rows=blank_rows,
        duplicates_removed=duplicates_removed,
        invalid=invalid,
        started_at=started_at,
        finished_at=finished_at,
        service_url=service_url,
    )


def format_summary(summary: RunSummary, report_path: Path | None = None) -> str:
    """Texto del resumen que se imprime al terminar."""
    lines = [
        "=" * 46,
        " RESUMEN DE LA EJECUCIÓN",
        "=" * 46,
        f" Total (cédulas únicas):  {summary.total}",
        f" Limpios:                 {summary.clean}",
        f" Con antecedentes:        {summary.with_records}",
        f" No encontrados:          {summary.not_found}",
        f" Errores:                 {summary.errors}"
        f"  (de ellos, {summary.invalid} por cédula inválida)",
        "-" * 46,
        f" Filas leídas:            {summary.rows_read}",
        f" Duplicadas eliminadas:   {summary.duplicates_removed}",
        f" Filas vacías ignoradas:  {summary.blank_rows}",
        f" Duración:                {summary.duration_seconds:.1f} s",
    ]
    if report_path is not None:
        lines.append(f" Reporte:                 {report_path}")
    lines.append("=" * 46)
    return "\n".join(lines)


def write_report(results: Sequence[QueryResult], summary: RunSummary, path: Path | str) -> Path:
    """Escribe el reporte en `path` de forma atómica (nunca deja un archivo a medias)."""
    target = Path(path)
    workbook = Workbook()

    sheet = workbook.active
    if not isinstance(sheet, Worksheet):  # pragma: no cover - un libro nuevo siempre trae hoja
        raise ReportWriteError("No se pudo crear la hoja del reporte")
    sheet.title = "Reporte"
    sheet.append(list(HEADERS))
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(vertical="center")

    for row_number, result in enumerate(results, start=2):
        sheet.append(
            [result.cedula, result.names, result.status.value, result.detail, result.queried_at]
        )
        # La cédula va como texto: no pierde ceros a la izquierda ni pasa a notación científica.
        sheet.cell(row=row_number, column=1).number_format = "@"
        sheet.cell(row=row_number, column=5).number_format = DATETIME_FORMAT
        sheet.cell(row=row_number, column=3).fill = _STATUS_FILL[result.status]

    for index, width in enumerate((16, 34, 18, 62, 20), start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions

    _write_summary_sheet(workbook.create_sheet("Resumen"), summary)

    try:
        target.parent.mkdir(parents=True, exist_ok=True)
        fd, tmp_name = tempfile.mkstemp(dir=target.parent, prefix=f".{target.stem}-", suffix=".tmp")
        os.close(fd)
        tmp = Path(tmp_name)
        try:
            workbook.save(tmp)
            os.replace(tmp, target)
        except BaseException:
            with contextlib.suppress(OSError):
                tmp.unlink()
            raise
    except OSError as exc:
        raise ReportWriteError(
            f"No se pudo escribir el reporte en '{target}': {exc}. "
            "Si el archivo está abierto en Excel, ciérrelo e intente de nuevo."
        ) from exc

    log.info("Reporte escrito en %s (%d filas)", target, len(results))
    return target


def _write_summary_sheet(sheet: Worksheet, summary: RunSummary) -> None:
    rows: list[tuple[str, object]] = [
        ("Total (cédulas únicas)", summary.total),
        ("Limpios", summary.clean),
        ("Con antecedentes", summary.with_records),
        ("No encontrados", summary.not_found),
        ("Errores", summary.errors),
        ("  de ellos, por cédula inválida", summary.invalid),
        ("Filas leídas", summary.rows_read),
        ("Duplicadas eliminadas", summary.duplicates_removed),
        ("Filas vacías ignoradas", summary.blank_rows),
        ("Inicio", summary.started_at),
        ("Fin", summary.finished_at),
        ("Duración (s)", round(summary.duration_seconds, 1)),
        ("Servicio consultado", summary.service_url),
    ]
    sheet.append(["Concepto", "Valor"])
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = _HEADER_FILL
    for label, value in rows:
        sheet.append([label, value])
    for row in sheet.iter_rows(min_row=2, min_col=2, max_col=2):
        for cell in row:
            if isinstance(cell.value, datetime):
                cell.number_format = DATETIME_FORMAT
            cell.alignment = Alignment(horizontal="left")
    sheet.column_dimensions["A"].width = 34
    sheet.column_dimensions["B"].width = 40
