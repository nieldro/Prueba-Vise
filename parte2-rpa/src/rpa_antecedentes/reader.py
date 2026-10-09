"""Lectura del Excel de entrada."""

from __future__ import annotations

import logging
import unicodedata
from dataclasses import dataclass
from pathlib import Path
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from rpa_antecedentes.models import InputRow

log = logging.getLogger(__name__)

SUPPORTED_SUFFIXES = frozenset({".xlsx", ".xlsm"})


class InputFileError(Exception):
    """El archivo de entrada no existe, no se puede abrir o no tiene la columna esperada."""


@dataclass(frozen=True, slots=True)
class ExcelRead:
    rows: list[InputRow]
    """Filas con contenido, en el orden del archivo."""
    blank_rows: int
    """Filas de datos cuya celda de cédula estaba vacía (se ignoran)."""


def _normalize_header(value: object) -> str:
    """Minúsculas, sin espacios ni tildes: 'Cédula ' -> 'cedula'."""
    text = unicodedata.normalize("NFD", str(value).strip().casefold())
    return "".join(ch for ch in text if not unicodedata.combining(ch))


def _cell_to_text(value: object) -> str:
    """Convierte una celda a texto. Excel guarda los números como `int` o `float` (123.0)."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def read_cedulas(path: Path | str, column: str = "cedula") -> ExcelRead:
    """Lee la columna `column` del primer libro/hoja activa del Excel.

    La primera fila con contenido se toma como encabezado; el nombre de la columna se compara
    sin distinguir mayúsculas, espacios ni tildes.

    Raises:
        InputFileError: si el archivo no existe, no es un .xlsx válido o no tiene la columna.
    """
    file = Path(path)
    if not file.is_file():
        raise InputFileError(f"No existe el archivo de entrada: {file}")
    if file.suffix.lower() not in SUPPORTED_SUFFIXES:
        raise InputFileError(
            f"Formato no soportado '{file.suffix}': se esperaba un archivo .xlsx ({file.name})"
        )

    try:
        workbook = load_workbook(file, read_only=True, data_only=True)
    except (InvalidFileException, BadZipFile, KeyError, OSError) as exc:
        raise InputFileError(f"No se pudo abrir el Excel '{file.name}': {exc}") from exc

    try:
        sheet = workbook.active
        if sheet is None:
            raise InputFileError(f"El Excel '{file.name}' no tiene hojas")

        rows: list[InputRow] = []
        blank = 0
        column_index: int | None = None
        wanted = _normalize_header(column)

        for row_number, values in enumerate(sheet.iter_rows(values_only=True), start=1):
            if column_index is None:
                if not any(cell is not None and str(cell).strip() for cell in values):
                    continue  # filas vacías antes del encabezado
                headers = [_normalize_header(cell) if cell is not None else "" for cell in values]
                if wanted not in headers:
                    found = ", ".join(str(c) for c in values if c is not None) or "(ninguna)"
                    raise InputFileError(
                        f"El Excel no tiene una columna llamada '{column}'. "
                        f"Columnas encontradas: {found}"
                    )
                column_index = headers.index(wanted)
                continue

            text = _cell_to_text(values[column_index]) if column_index < len(values) else ""
            if text:
                rows.append(InputRow(row=row_number, raw=text))
            else:
                blank += 1

        if column_index is None:
            raise InputFileError(f"El Excel '{file.name}' está vacío")
    finally:
        workbook.close()

    log.info("Leídas %d filas con cédula de %s (%d vacías ignoradas)", len(rows), file.name, blank)
    return ExcelRead(rows=rows, blank_rows=blank)
