from __future__ import annotations

from datetime import datetime, timedelta
from pathlib import Path

import pytest
from openpyxl import load_workbook

from rpa_antecedentes.models import QueryResult, Status
from rpa_antecedentes.report import (
    HEADERS,
    ReportWriteError,
    build_summary,
    format_summary,
    write_report,
)

START = datetime(2026, 10, 9, 10, 0, 0)
END = START + timedelta(seconds=12)
WHEN = datetime(2026, 10, 9, 10, 0, 5)


def result(
    cedula: str, status: Status, *, names: str = "", detail: str = "", when: datetime | None = WHEN
) -> QueryResult:
    return QueryResult(cedula=cedula, names=names, status=status, detail=detail, queried_at=when)


@pytest.fixture
def results() -> list[QueryResult]:
    return [
        result("0012345", Status.CLEAN, names="ANA PÉREZ", detail="No registra antecedentes"),
        result("2222222", Status.WITH_RECORDS, names="LUIS GÓMEZ", detail="Registra 1 antecedente"),
        result("3333333", Status.NOT_FOUND, detail="La cédula no figura en el servicio"),
        result("4444444", Status.ERROR, detail="El servicio no respondió"),
        result("12", Status.ERROR, detail="Cédula inválida: ...", when=None),
    ]


def summary_for(results: list[QueryResult]):
    return build_summary(
        results,
        rows_read=8,
        blank_rows=1,
        duplicates_removed=2,
        started_at=START,
        finished_at=END,
        service_url="http://servicio.test/api/v1/antecedentes",
    )


def test_summary_counts_each_status(results: list[QueryResult]) -> None:
    summary = summary_for(results)

    assert (summary.total, summary.clean, summary.with_records) == (5, 1, 1)
    assert (summary.not_found, summary.errors, summary.invalid) == (1, 2, 1)
    assert summary.duration_seconds == 12


def test_report_has_the_required_columns_in_order(
    results: list[QueryResult], tmp_path: Path
) -> None:
    path = write_report(results, summary_for(results), tmp_path / "reporte_antecedentes.xlsx")

    sheet = load_workbook(path)["Reporte"]
    assert tuple(cell.value for cell in sheet[1]) == HEADERS
    assert HEADERS == ("cedula", "nombres", "estado", "detalle", "fecha_consulta")


def test_report_rows_match_the_results(results: list[QueryResult], tmp_path: Path) -> None:
    path = write_report(results, summary_for(results), tmp_path / "r.xlsx")

    sheet = load_workbook(path)["Reporte"]
    rows = [tuple(c.value for c in row) for row in sheet.iter_rows(min_row=2)]

    assert len(rows) == 5
    assert rows[0] == ("0012345", "ANA PÉREZ", "limpio", "No registra antecedentes", WHEN)
    assert [r[2] for r in rows] == ["limpio", "con antecedentes", "no encontrado", "error", "error"]
    assert rows[4][4] is None  # la cédula inválida no se consultó: sin fecha


def test_cedulas_are_stored_as_text_keeping_leading_zeros(
    results: list[QueryResult], tmp_path: Path
) -> None:
    path = write_report(results, summary_for(results), tmp_path / "r.xlsx")

    cell = load_workbook(path)["Reporte"]["A2"]

    assert cell.value == "0012345"
    assert cell.data_type == "s"


def test_report_includes_a_summary_sheet(results: list[QueryResult], tmp_path: Path) -> None:
    path = write_report(results, summary_for(results), tmp_path / "r.xlsx")

    sheet = load_workbook(path)["Resumen"]
    values = {row[0].value: row[1].value for row in sheet.iter_rows(min_row=2)}

    assert values["Total (cédulas únicas)"] == 5
    assert values["Limpios"] == 1
    assert values["Duplicadas eliminadas"] == 2
    assert values["Servicio consultado"] == "http://servicio.test/api/v1/antecedentes"


def test_empty_results_still_produce_a_valid_report(tmp_path: Path) -> None:
    path = write_report([], summary_for([]), tmp_path / "vacio.xlsx")

    sheet = load_workbook(path)["Reporte"]
    assert sheet.max_row == 1
    assert tuple(cell.value for cell in sheet[1]) == HEADERS


def test_parent_directories_are_created(results: list[QueryResult], tmp_path: Path) -> None:
    path = write_report(results, summary_for(results), tmp_path / "a" / "b" / "r.xlsx")

    assert path.is_file()


def test_no_temporary_files_are_left_behind(results: list[QueryResult], tmp_path: Path) -> None:
    write_report(results, summary_for(results), tmp_path / "r.xlsx")

    assert [p.name for p in tmp_path.iterdir()] == ["r.xlsx"]


def test_existing_report_is_replaced(results: list[QueryResult], tmp_path: Path) -> None:
    target = tmp_path / "r.xlsx"
    write_report(results[:1], summary_for(results[:1]), target)
    write_report(results, summary_for(results), target)

    assert load_workbook(target)["Reporte"].max_row == 6


def test_write_failure_is_reported_clearly_and_cleans_up(
    results: list[QueryResult], tmp_path: Path
) -> None:
    blocker = tmp_path / "es_un_archivo"
    blocker.write_text("x", encoding="utf-8")

    with pytest.raises(ReportWriteError, match="No se pudo escribir"):
        write_report(results, summary_for(results), blocker / "r.xlsx")


def test_console_summary_shows_the_required_figures(
    results: list[QueryResult], tmp_path: Path
) -> None:
    text = format_summary(summary_for(results), tmp_path / "r.xlsx")

    for expected in (
        "Total (cédulas únicas):  5",
        "Limpios:                 1",
        "Con antecedentes:        1",
        "Errores:                 2",
    ):
        assert expected in text
    assert "Duplicadas eliminadas:   2" in text
    assert "1 por cédula inválida" in text
    assert "r.xlsx" in text
