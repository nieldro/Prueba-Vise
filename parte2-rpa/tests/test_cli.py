"""Pruebas de integración: el flujo completo contra el servicio simulado real."""

from __future__ import annotations

import logging
import os
import re
import socket
import subprocess
import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook

from rpa_antecedentes import cli
from rpa_antecedentes.config import parse_settings

from .conftest import ExcelFactory

FAST = ["--timeout", "0.5", "--retries", "2", "--backoff-base", "0.01", "--backoff-max", "0.05"]


def run_cli(excel: Path, out: Path, url: str, *extra: str) -> int:
    return cli.main(["-i", str(excel), "-o", str(out), "--url", url, *FAST, *extra])


def read_report(path: Path) -> list[tuple[object, ...]]:
    sheet = load_workbook(path)["Reporte"]
    return [tuple(c.value for c in row) for row in sheet.iter_rows(min_row=2)]


@pytest.fixture(autouse=True)
def _reset_logging() -> None:
    yield  # type: ignore[misc]
    logger = logging.getLogger("rpa_antecedentes")
    for handler in list(logger.handlers):
        logger.removeHandler(handler)
        handler.close()


def free_port() -> int:
    with socket.socket() as sock:
        sock.bind(("127.0.0.1", 0))
        return int(sock.getsockname()[1])


def test_full_run_produces_the_report_and_summary(
    make_excel: ExcelFactory, mock_service: str, tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    excel = make_excel(
        [
            "1012345678",  # limpio
            1098765409,  # con antecedentes (celda numérica)
            "63777103",  # no encontrado
            "1012345678",  # duplicada
            "  1012345678 ",  # duplicada con espacios
            "12345",  # inválida
            "ABC1234567",  # inválida
            None,  # vacía
            "55558888",  # transitoria: 503, 503, 200 (necesita 2 reintentos)
            "77779999",  # persistente: error
        ]
    )
    out = tmp_path / "salida" / "reporte_antecedentes.xlsx"

    code = run_cli(excel, out, mock_service)

    assert code == 0
    rows = read_report(out)
    by_cedula = {r[0]: r for r in rows}
    assert len(rows) == 7  # 9 con cédula - 2 duplicadas
    assert by_cedula["1012345678"][2] == "limpio"
    assert by_cedula["1098765409"][2] == "con antecedentes"
    assert by_cedula["63777103"][2] == "no encontrado"
    assert by_cedula["55558888"][2] == "limpio"  # se recuperó con reintentos
    assert by_cedula["77779999"][2] == "error"
    assert "4 intentos" in str(by_cedula["77779999"][3]) or "3 intentos" in str(
        by_cedula["77779999"][3]
    )
    assert by_cedula["12345"][2] == "error"
    assert "Cédula inválida" in str(by_cedula["12345"][3])
    assert by_cedula["12345"][4] is None

    printed = capsys.readouterr().out
    assert "Total (cédulas únicas):  7" in printed
    assert "Limpios:                 2" in printed
    assert "Con antecedentes:        1" in printed
    assert "Errores:                 3" in printed
    assert "Duplicadas eliminadas:   2" in printed


def test_a_failing_cedula_does_not_stop_the_process(
    make_excel: ExcelFactory, mock_service: str, tmp_path: Path
) -> None:
    excel = make_excel(["77779999", "1012345678"])  # la primera falla siempre
    out = tmp_path / "r.xlsx"

    assert run_cli(excel, out, mock_service) == 0

    assert [r[2] for r in read_report(out)] == ["error", "limpio"]


def test_service_down_marks_everything_as_error_but_still_reports(
    make_excel: ExcelFactory, tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    excel = make_excel(["1012345678", "1098765409", "63777103"])
    out = tmp_path / "r.xlsx"

    code = run_cli(excel, out, f"http://127.0.0.1:{free_port()}")  # nadie escucha ahí

    assert code == 0
    rows = read_report(out)
    assert [r[2] for r in rows] == ["error", "error", "error"]
    assert all("no respondió" in str(r[3]) for r in rows)
    assert "Errores:                 3" in capsys.readouterr().out


def test_workers_give_the_same_report_as_sequential(
    make_excel: ExcelFactory, mock_service: str, tmp_path: Path
) -> None:
    values = [f"{n:08d}" for n in range(10000100, 10000130)]
    excel = make_excel(values)

    run_cli(excel, tmp_path / "uno.xlsx", mock_service)
    run_cli(excel, tmp_path / "cuatro.xlsx", mock_service, "--workers", "4")

    project = lambda rows: [(r[0], r[1], r[2], r[3]) for r in rows]  # noqa: E731
    assert project(read_report(tmp_path / "uno.xlsx")) == project(
        read_report(tmp_path / "cuatro.xlsx")
    )


def test_log_file_has_timestamps_and_masked_ids(
    make_excel: ExcelFactory, mock_service: str, tmp_path: Path
) -> None:
    excel = make_excel(["1012345678", "77779999"])
    out = tmp_path / "r.xlsx"

    run_cli(excel, out, mock_service)

    log_text = (tmp_path / "rpa_antecedentes.log").read_text(encoding="utf-8")
    assert re.search(r"^\d{4}-\d{2}-\d{2} \d{2}:\d{2}:\d{2} \| INFO", log_text, re.MULTILINE)
    assert "consulta cedula=*******678 intento=1/3 http=200" in log_text
    assert "WARNING" in log_text
    assert "1012345678" not in log_text


def test_full_ids_in_log_when_requested(
    make_excel: ExcelFactory, mock_service: str, tmp_path: Path
) -> None:
    run_cli(make_excel(["1012345678"]), tmp_path / "r.xlsx", mock_service, "--log-full-ids")

    assert "cedula=1012345678" in (tmp_path / "rpa_antecedentes.log").read_text(encoding="utf-8")


def test_custom_log_file(make_excel: ExcelFactory, mock_service: str, tmp_path: Path) -> None:
    log_path = tmp_path / "logs" / "mi.log"

    run_cli(
        make_excel(["1012345678"]), tmp_path / "r.xlsx", mock_service, "--log-file", str(log_path)
    )

    assert log_path.is_file()


def test_empty_excel_produces_an_empty_report(
    make_excel: ExcelFactory, mock_service: str, tmp_path: Path
) -> None:
    out = tmp_path / "r.xlsx"

    assert run_cli(make_excel([]), out, mock_service) == 0
    assert read_report(out) == []


def test_missing_input_file_exits_with_1(
    tmp_path: Path, mock_service: str, capsys: pytest.CaptureFixture[str]
) -> None:
    code = run_cli(tmp_path / "no_existe.xlsx", tmp_path / "r.xlsx", mock_service)

    assert code == 1
    assert "No existe el archivo" in capsys.readouterr().err
    assert not (tmp_path / "r.xlsx").exists()


def test_input_without_cedula_column_exits_with_1(
    make_excel: ExcelFactory, mock_service: str, tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    excel = make_excel(["1234567"], header="documento")

    assert run_cli(excel, tmp_path / "r.xlsx", mock_service) == 1
    assert "columna llamada 'cedula'" in capsys.readouterr().err


def test_real_portal_url_exits_with_2_and_explains(
    make_excel: ExcelFactory, tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    code = run_cli(
        make_excel(["1234567"]),
        tmp_path / "r.xlsx",
        "https://apps.procuraduria.gov.co/webcert/inicio.aspx?tpo=2",
    )

    assert code == 2
    assert "autorización" in capsys.readouterr().err


def test_report_that_cannot_be_written_exits_with_1(
    make_excel: ExcelFactory, mock_service: str, tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    blocker = tmp_path / "archivo"
    blocker.write_text("x", encoding="utf-8")

    code = run_cli(make_excel(["1012345678"]), blocker / "r.xlsx", mock_service)

    assert code == 1
    assert "Error de salida" in capsys.readouterr().err


def test_keyboard_interrupt_exits_with_130(
    monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str]
) -> None:
    def interrupted(_settings: object) -> int:
        raise KeyboardInterrupt

    monkeypatch.setattr(cli, "run", interrupted)

    assert cli.main(["-i", "x.xlsx"]) == 130
    assert "interrumpido" in capsys.readouterr().err


# --- parámetros -------------------------------------------------------------------------------


def test_defaults_and_environment_variables() -> None:
    settings = parse_settings(
        [],
        {
            "ANTECEDENTES_URL": "http://servicio:9000",
            "ANTECEDENTES_TIMEOUT": "2.5",
            "ANTECEDENTES_RETRIES": "5",
            "RPA_WORKERS": "3",
            "RPA_OUTPUT": "/datos/salida/reporte.xlsx",
            "RPA_LOG_LEVEL": "debug",
        },
    )

    assert settings.service_url == "http://servicio:9000"
    assert settings.timeout == 2.5
    assert settings.max_retries == 5
    assert settings.workers == 3
    assert settings.log_level == "DEBUG"
    assert settings.log_file == Path("/datos/salida/rpa_antecedentes.log")
    assert settings.mask_ids is True


def test_command_line_overrides_the_environment() -> None:
    settings = parse_settings(
        ["--url", "http://otro:1", "--workers", "2"],
        {"ANTECEDENTES_URL": "http://env:9", "RPA_WORKERS": "8"},
    )

    assert (settings.service_url, settings.workers) == ("http://otro:1", 2)


def test_log_full_ids_from_environment() -> None:
    assert parse_settings([], {"RPA_LOG_FULL_IDS": "true"}).mask_ids is False


@pytest.mark.parametrize(
    "argv",
    [
        ["--workers", "0"],
        ["--workers", "99"],
        ["--retries", "-1"],
        ["--log-level", "NOISY"],
        ["--timeout", "abc"],
    ],
)
def test_invalid_arguments_exit_with_2(argv: list[str], capsys: pytest.CaptureFixture[str]) -> None:
    with pytest.raises(SystemExit) as exit_info:
        parse_settings(argv, {})

    assert exit_info.value.code == 2
    assert capsys.readouterr().err


def test_invalid_numeric_environment_variable_is_reported() -> None:
    with pytest.raises(SystemExit, match="ANTECEDENTES_TIMEOUT"):
        parse_settings([], {"ANTECEDENTES_TIMEOUT": "rápido"})


def test_help_describes_every_option(capsys: pytest.CaptureFixture[str]) -> None:
    with pytest.raises(SystemExit) as exit_info:
        parse_settings(["--help"], {})

    assert exit_info.value.code == 0
    help_text = capsys.readouterr().out
    for option in (
        "--input",
        "--output",
        "--url",
        "--timeout",
        "--retries",
        "--workers",
        "--log-file",
    ):
        assert option in help_text


def test_module_entry_point_runs() -> None:
    env = {**os.environ, "PYTHONPATH": str(Path(__file__).resolve().parents[1] / "src")}
    completed = subprocess.run(
        [sys.executable, "-m", "rpa_antecedentes", "--help"],
        capture_output=True,
        text=True,
        env=env,
        timeout=30,
        check=False,
    )

    assert completed.returncode == 0
    assert "usage: rpa-antecedentes" in completed.stdout
